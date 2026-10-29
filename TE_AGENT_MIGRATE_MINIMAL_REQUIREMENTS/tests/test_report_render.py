import io
import os
import re
import shutil
import sys
import tempfile
import time
import unittest
from unittest import mock

try:
    from tests.test_report_fixtures import SECRET_MARKER, build_discovery, write_selection
except ImportError:
    from test_report_fixtures import SECRET_MARKER, build_discovery, write_selection

from te_report import render_html, render_xlsx
from te_report.__main__ import main
from te_report.graph import Graph, build_graph
from te_report.impact import flag_impact, resolve_selection
from te_report.loader import load_discovery

try:
    import openpyxl  # noqa: F401
    HAVE_OPENPYXL = True
except ImportError:
    HAVE_OPENPYXL = False

EVIL = "<img src=x onerror=alert(1)>"


class RenderBase(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        self.root = os.path.join(self.tmp, "discovery_output")
        build_discovery(self.root)
        self.graph = build_graph(load_discovery(self.root))
        self.res = resolve_selection(self.graph, ["agent-a", "10.0.1.2", "agent-off"])
        self.findings = flag_impact(self.graph, self.res["selected"])

    def tearDown(self):
        shutil.rmtree(self.tmp)

    def page(self, selection=True):
        if selection:
            return render_html.render(self.graph, self.findings, self.res, {"generated": "2026-10-28 12:00"})
        return render_html.render(self.graph, [], None, {})


class HtmlTests(RenderBase):
    def test_self_contained_single_file(self):
        text = self.page()
        self.assertTrue(text.startswith("<!DOCTYPE html>"))
        self.assertNotIn("<link", text)
        self.assertNotIn("@import", text)
        self.assertNotIn("url(", text)
        self.assertNotIn("<img", text)
        self.assertNotIn(" src=", text)
        self.assertNotIn("http://", text)
        self.assertNotIn("https://", text)
        self.assertEqual(text.count("<script>"), 1)
        self.assertIn("Content-Security-Policy", text)
        for bad in ("eval(", "innerHTML", "outerHTML", "document.write", "new Function"):
            self.assertNotIn(bad, text)

    def test_sections_and_badges(self):
        text = self.page()
        for heading in ("1. Summary", "2. Inventory tree", "3. Findings", "4. Offline agents", "5. Shares"):
            self.assertIn(heading, text)
        for label in ("BREAKS", "REVIEW", "OK"):
            self.assertRegex(text, r'<span class="b b-%s">%s</span>' % (label, label))
        self.assertIn("SELECTED", text)
        self.assertIn('id="fbox"', text)
        self.assertIn('data-level="BREAKS"', text)
        # selected agents and the groups they are shared with
        self.assertIn("Source", text)
        self.assertIn("Other", text)

    def test_findings_sorted_worst_first(self):
        text = self.page()
        levels = re.findall(r'<tr data-level="(\w+)">', text)
        self.assertTrue(levels)
        order = {"BREAKS": 0, "REVIEW": 1, "OK": 2}
        ranks = [order[x] for x in levels]
        self.assertEqual(ranks, sorted(ranks))
        self.assertEqual(len(levels), len(self.findings))

    def test_tree_collapsed_except_top_level(self):
        text = self.page()
        tree = text[text.index('<div id="tree">'):text.index('id="findings"')]
        opened = re.findall(r"<details[^>]* open", tree)
        groups = len(self.graph.nodes_of("group"))
        self.assertEqual(len(opened), groups)
        self.assertTrue(all('class="top"' in o for o in opened))
        self.assertGreater(tree.count("<details"), groups)

    def test_tree_has_kind_name_id_per_node(self):
        text = self.page()
        self.assertIn('<span class="kind">agent</span> <span class="name">agent-a</span> <span class="id">1</span>', text)
        self.assertIn('<span class="kind">cluster</span> <span class="name">cluster-c</span>', text)
        self.assertIn('<span class="kind">widget</span> <span class="name">agent-id widget</span>', text)
        self.assertIn('<span class="kind">member</span> <span class="name">node1</span>', text)

    def test_ids_unique_and_links_resolve(self):
        text = self.page()
        ids = re.findall(r'\sid="([^"]+)"', text)
        self.assertEqual(len(ids), len(set(ids)))
        for target in re.findall(r'href="#([^"]+)"', text):
            self.assertIn(target, ids)
        # every graph test, dashboard and widget is drawn in full
        for kind in ("test", "dashboard", "widget", "agent", "cluster"):
            for node in self.graph.nodes_of(kind):
                self.assertIn(">%s</span>" % render_html.esc(node["name"]), text)

    def test_shared_test_drawn_once(self):
        text = self.page()
        full = re.findall(r'<(?:div class="leaf"|details) id="n\d+">(?:<summary>)?<span class="kind">test</span> '
                          r'<span class="name">test-two</span>', text)
        self.assertEqual(len(full), 1)  # assigned to agent-a and agent-b
        self.assertIn('<div class="ref"><span class="kind">test</span> <a class="goto"', text)

    def test_no_secrets(self):
        text = self.page() + self.page(selection=False)
        self.assertNotIn(SECRET_MARKER, text)
        self.assertNotIn("transactionScript", text)
        self.assertNotIn("agentToken", text)

    def test_defensive_filter_on_node_attrs(self):
        g = Graph()
        g.add_node("group:1", "group", "G", aid="1")
        g.add_node(
            "agent:9", "agent", "a9", agent_id="9", state="online", groups=["1"],
            password=SECRET_MARKER, agentToken=SECRET_MARKER, ips=["10.9.9.9"],
        )
        g.add_node(
            "test:7", "test", "t7", test_id="7", test_type="http-server", aid="1", aids=["1"],
            transactionScript=SECRET_MARKER, client_secret=SECRET_MARKER,
        )
        g.add_edge("agent:9", "group:1", "visible_in")
        g.add_edge("agent:9", "test:7", "assigned_test", group="1")
        text = render_html.render(g, [], None, {})
        self.assertNotIn(SECRET_MARKER, text)
        self.assertIn("a9", text)
        cleaned = render_html.clean_attrs({"state": "x", "Password": "p", "x_token": "t", "n": {"a": 1}, "l": [1, {"a": 2}]})
        self.assertEqual(cleaned, {"state": "x", "l": [1]})

    def test_html_escaping_everywhere(self):
        g = Graph()
        g.add_node("group:1", "group", EVIL, aid="1")
        g.add_node("agent:9", "agent", EVIL, agent_id=EVIL, state="online", groups=["1"], ips=[EVIL], shared_with=["1", EVIL])
        g.add_node("test:7", "test", EVIL, test_id="7", test_type=EVIL, aid="1", aids=["1"])
        g.add_node("dashboard:1:d", "dashboard", EVIL, dashboard_id="d", aid="1")
        g.add_node("widget:1:d:w", "widget", EVIL, widget_id="w", aid="1", widget_type=EVIL, filter_keys=[EVIL])
        g.add_edge("agent:9", "group:1", "visible_in")
        g.add_edge("agent:9", "test:7", "assigned_test", group="1")
        g.add_edge("dashboard:1:d", "widget:1:d:w", "has_widget")
        g.add_edge("widget:1:d:w", "test:7", "filters_on_test", via="TEST", mode="id")
        res = {"selected": set(["agent:9"]), "excluded_offline": set(), "unmatched": [EVIL], "matches": {}}
        findings = [{
            "level": "BREAKS", "element_kind": EVIL, "element_id": "agent:9", "element_name": EVIL,
            "reason": EVIL, "fix": EVIL, "fix_owner": EVIL, "rule": "r", "group": "1",
        }]
        text = render_html.render(g, findings, res, {"input": EVIL, "generated": EVIL})
        self.assertNotIn(EVIL, text)
        self.assertNotIn("<img", text)
        self.assertIn("&lt;img src=x onerror=alert(1)&gt;", text)

    def test_inventory_only(self):
        text = self.page(selection=False)
        self.assertIn("No selection was given", text)
        self.assertIn("4. Offline agents", text)
        self.assertNotIn('id="ftable"', text)

    def test_offline_list(self):
        text = self.page()
        section = text[text.index('id="offline"'):text.index('id="shares"')]
        self.assertIn("agent-off", section)
        self.assertIn("cluster-bad", section)  # in the all-offline details
        self.assertIn("agent-disabled", section)

    def test_large_graph_stays_small_and_fast(self):
        g = Graph()
        g.add_node("group:1", "group", "Big", aid="1")
        for t in range(300):
            g.add_node("test:%d" % t, "test", "t%d" % t, test_id=str(t), aid="1", aids=["1"])
        for a in range(200):
            aid = "agent:%d" % a
            g.add_node(aid, "agent", "a%d" % a, agent_id=str(a), state="online", groups=["1"])
            g.add_edge(aid, "group:1", "visible_in")
            for t in range(300):
                g.add_edge(aid, "test:%d" % t, "assigned_test", group="1")
        for d in range(20):
            g.add_node("dashboard:1:%d" % d, "dashboard", "d%d" % d, dashboard_id=str(d), aid="1")
            for w in range(50):
                wid = "widget:1:%d:%d" % (d, w)
                g.add_node(wid, "widget", "w%d" % w, widget_id=str(w), aid="1", widget_type="Table")
                g.add_edge("dashboard:1:%d" % d, wid, "has_widget")
                for t in range(0, 300, 5):
                    g.add_edge(wid, "test:%d" % t, "filters_on_test", via="TEST", mode="id")
        start = time.time()
        text = render_html.render(g, [], None, {})
        self.assertLess(time.time() - start, 30)
        # 300 tests are drawn in full once; 60,000 agent->test pairs are one-line references
        full_tests = len(re.findall(r'id="n\d+">(?:<summary>)?<span class="kind">test</span>', text))
        self.assertEqual(full_tests, 300)
        self.assertLess(len(text.encode("utf-8")), 15 * 1024 * 1024)

    def test_cli_html(self):
        out = os.path.join(self.tmp, "out")
        sel = write_selection(os.path.join(self.tmp, "sel.txt"), ["agent-a"])
        buf = io.StringIO()
        self.assertEqual(main(["--input", self.root, "--output", out, "--selection", sel, "--html"], out=buf), 0)
        path = os.path.join(out, "report.html")
        self.assertTrue(os.path.isfile(path))
        self.assertIn("report.html", buf.getvalue())
        self.assertNotIn(SECRET_MARKER, buf.getvalue())
        with io.open(path, encoding="utf-8") as handle:
            body = handle.read()
        self.assertIn("3. Findings", body)
        self.assertNotIn(SECRET_MARKER, body)
        self.assertFalse(os.path.exists(os.path.join(out, "report.xlsx")))

    def test_cli_html_without_selection(self):
        out = os.path.join(self.tmp, "out")
        self.assertEqual(main(["--input", self.root, "--output", out, "--html"], out=io.StringIO()), 0)
        self.assertTrue(os.path.isfile(os.path.join(out, "report.html")))


class XlsxMissingTests(RenderBase):
    def test_missing_openpyxl_message(self):
        with mock.patch.dict(sys.modules, {"openpyxl": None}):
            with self.assertRaises(RuntimeError) as ctx:
                render_xlsx.write_workbook(os.path.join(self.tmp, "x.xlsx"), self.graph, [], None, {})
        self.assertEqual(str(ctx.exception), "openpyxl is not installed; use --html or the CSV files")
        self.assertFalse(os.path.exists(os.path.join(self.tmp, "x.xlsx")))

    def test_cli_xlsx_missing_openpyxl_writes_others_then_fails(self):
        out = os.path.join(self.tmp, "out")
        sel = write_selection(os.path.join(self.tmp, "sel.txt"), ["agent-a"])
        buf = io.StringIO()
        with mock.patch.dict(sys.modules, {"openpyxl": None}):
            code = main(["--input", self.root, "--output", out, "--selection", sel, "--html", "--xlsx"], out=buf)
        self.assertEqual(code, 1)
        self.assertIn("openpyxl is not installed; use --html or the CSV files", buf.getvalue())
        self.assertTrue(os.path.isfile(os.path.join(out, "report.html")))
        self.assertTrue(os.path.isfile(os.path.join(out, "findings.csv")))
        self.assertFalse(os.path.exists(os.path.join(out, "report.xlsx")))
        self.assertNotIn(SECRET_MARKER, buf.getvalue())


@unittest.skipUnless(HAVE_OPENPYXL, "openpyxl is not installed")
class XlsxTests(RenderBase):
    def write(self, graph=None, findings=None, res="default"):
        path = os.path.join(self.tmp, "r.xlsx")
        render_xlsx.write_workbook(
            path, graph or self.graph, self.findings if findings is None else findings,
            self.res if res == "default" else res, {"generated": "now", "input": "in"},
        )
        return openpyxl.load_workbook(path)

    def test_sheets_and_headers(self):
        wb = self.write()
        self.assertEqual(wb.sheetnames, [
            "Summary", "Selected", "Offline", "Shares", "Tree", "Findings", "Tests", "Dashboards", "Alerts"])
        for name in wb.sheetnames:
            sheet = wb[name]
            self.assertEqual(sheet.freeze_panes, "A2", name)
            self.assertTrue(sheet.auto_filter.ref.startswith("A1:"), name)
            self.assertTrue(sheet["A1"].value, name)
            self.assertTrue(sheet.column_dimensions["A"].width, name)
        findings = wb["Findings"]
        self.assertEqual(findings.max_row, len(self.findings) + 1)
        self.assertEqual(findings["A2"].value, "BREAKS")
        selected = wb["Selected"]
        self.assertEqual(sorted(r[0].value for r in selected.iter_rows(min_row=2)), ["1", "4"])
        self.assertEqual([r[0].value for r in wb["Offline"].iter_rows(min_row=2)], ["3"])
        summary = dict((r[0].value, r[1].value) for r in wb["Summary"].iter_rows(min_row=2))
        self.assertEqual(summary["Selected agents/clusters"], 2)
        tests = wb["Tests"]
        self.assertEqual(tests.max_row, len(self.graph.nodes_of("test")) + 1)
        self.assertGreater(wb["Dashboards"].max_row, 1)
        self.assertGreater(wb["Alerts"].max_row, 1)
        self.assertGreater(wb["Tree"].max_row, 10)
        self.assertGreater(wb["Shares"].max_row, 1)

    def test_inventory_only_has_empty_sheets_with_headers(self):
        wb = self.write(findings=[], res=None)
        self.assertEqual(wb["Findings"].max_row, 1)
        self.assertEqual(wb["Selected"].max_row, 1)
        self.assertEqual(wb["Findings"]["A1"].value, "level")
        self.assertEqual(sorted(r[0].value for r in wb["Offline"].iter_rows(min_row=2)), ["3", "5", "6"])

    def test_formula_injection_neutralised(self):
        g = Graph()
        g.add_node("group:1", "group", "=cmd()", aid="1")
        g.add_node("agent:9", "agent", "@SUM(A1)", agent_id="9", state="online", groups=["1"], shared_count=1)
        g.add_node("test:7", "test", "+evil", test_id="7", aid="1", aids=["1"], test_type="-x")
        g.add_edge("agent:9", "group:1", "visible_in")
        g.add_edge("agent:9", "test:7", "assigned_test", group="1")
        res = {"selected": set(["agent:9"]), "excluded_offline": set(), "unmatched": [], "matches": {}}
        findings = [{
            "level": "REVIEW", "element_kind": "agent", "element_id": "agent:9", "element_name": "=1+1",
            "reason": "-2+3", "fix": "@x", "fix_owner": "tool", "rule": "r", "group": "1",
        }]
        wb = self.write(g, findings, res)
        seen = 0
        for sheet in wb.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    value = cell.value
                    if isinstance(value, str) and value[:1] in "=+-@":
                        self.assertRegex(value, r"^-?[0-9][0-9.,]*$", "%s!%s" % (sheet.title, cell.coordinate))
                    if isinstance(value, str) and value.startswith("'"):
                        seen += 1
                    self.assertNotEqual(cell.data_type, "f", "%s!%s" % (sheet.title, cell.coordinate))
        self.assertGreaterEqual(seen, 5)
        self.assertEqual(wb["Findings"]["D2"].value, "'=1+1")

    def test_no_secrets_in_workbook(self):
        wb = self.write()
        for sheet in wb.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    text = str(cell.value)
                    self.assertNotIn(SECRET_MARKER, text)
                    self.assertNotIn("transactionScript", text)

    def test_cli_xlsx(self):
        out = os.path.join(self.tmp, "out")
        sel = write_selection(os.path.join(self.tmp, "sel.txt"), ["agent-a"])
        self.assertEqual(
            main(["--input", self.root, "--output", out, "--selection", sel, "--html", "--xlsx"], out=io.StringIO()), 0)
        self.assertTrue(os.path.isfile(os.path.join(out, "report.xlsx")))
        self.assertTrue(os.path.isfile(os.path.join(out, "report.html")))


if __name__ == "__main__":
    unittest.main()
