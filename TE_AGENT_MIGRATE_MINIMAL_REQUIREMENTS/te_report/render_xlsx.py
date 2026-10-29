"""Excel workbook for the report. ``openpyxl`` is imported inside the writer only."""

import re

from te_report.impact import LEVELS, is_offline, summarize
from te_report.outputs import (
    AGENT_COLUMNS, FINDING_COLUMNS, SHARE_COLUMNS, TREE_COLUMNS,
    all_offline_ids, offline_rows, safe_cell, selected_rows, share_rows, tree_rows,
    _matched_by,
)
from te_report.render_html import clean_attrs, worst_levels

MISSING_MESSAGE = "openpyxl is not installed; use --html or the CSV files"
MAX_CELL = 32000
MAX_ROWS = 1000000  # Excel limit is 1,048,576 rows per sheet
SHEETS = (
    "Summary", "Selected", "Offline", "Shares", "Tree", "Findings", "Tests", "Dashboards", "Alerts",
)

_ILLEGAL = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f]")

TEST_COLUMNS = [
    "test_id", "name", "type", "enabled", "groups", "agents", "agent_count",
    "alert_rules", "tags", "flag", "reason",
]
DASHBOARD_COLUMNS = [
    "group", "dashboard_id", "dashboard", "widget_id", "widget", "widget_type",
    "filters", "agents_hit", "selected_agents_hit", "flag", "reason",
]
ALERT_COLUMNS = [
    "kind", "id", "name", "group", "detail", "tests", "agents", "flag", "reason",
]


def _cell(value):
    """Cell value: numbers stay numbers, text is made formula-safe and clean."""
    if isinstance(value, bool) or value is None:
        return safe_cell(value)
    if isinstance(value, (int, float)):
        return value
    text = _ILLEGAL.sub("", safe_cell(value))
    if len(text) > MAX_CELL:
        text = text[:MAX_CELL] + "..."
    return text


def _names(graph, ids, limit=8):
    names = [graph.nodes[i]["name"] or i for i in sorted(ids, key=lambda i: ((graph.nodes[i]["name"] or "").lower(), i))]
    text = "; ".join(names[:limit])
    if len(names) > limit:
        text += " (+%d more)" % (len(names) - limit)
    return text


def _group_names(graph, aids):
    out = []
    for aid in aids:
        node = graph.nodes.get("group:%s" % aid)
        out.append((node["name"] if node else None) or str(aid))
    return "; ".join(out)


def _level_info(findings):
    """element id -> (worst level, first reason of that level)."""
    worst = worst_levels(findings)
    reason = {}
    for item in findings:
        if worst.get(item["element_id"]) == item["level"] and item["element_id"] not in reason:
            reason[item["element_id"]] = item["reason"]
    return worst, reason


def tests_rows(graph, findings):
    worst, reason = _level_info(findings)
    rows = []
    for node in graph.nodes_of("test"):
        attrs = clean_attrs(node["attrs"])
        agents = graph.parents(node["id"], "assigned_test")
        rules = graph.children(node["id"], "alert_rule")
        tags = graph.children(node["id"], "tagged")
        rows.append({
            "test_id": attrs.get("test_id"),
            "name": node["name"],
            "type": attrs.get("test_type") or "",
            "enabled": attrs.get("enabled"),
            "groups": _group_names(graph, attrs.get("aids") or []),
            "agents": _names(graph, agents),
            "agent_count": len(agents),
            "alert_rules": _names(graph, rules),
            "tags": _names(graph, tags),
            "flag": worst.get(node["id"], ""),
            "reason": reason.get(node["id"], ""),
        })
    rows.sort(key=lambda r: ((r["name"] or "").lower(), str(r["test_id"])))
    return rows


def dashboard_rows(graph, findings, selected):
    worst, reason = _level_info(findings)
    rows = []
    dashboards = sorted(graph.nodes_of("dashboard"), key=lambda n: ((n["name"] or "").lower(), n["id"]))
    for dash in dashboards:
        attrs = clean_attrs(dash["attrs"])
        group = _group_names(graph, [attrs.get("aid")])
        widgets = sorted(graph.children(dash["id"], "has_widget"),
                         key=lambda i: ((graph.nodes[i]["name"] or "").lower(), i))
        if not widgets:
            rows.append({
                "group": group, "dashboard_id": attrs.get("dashboard_id"), "dashboard": dash["name"],
                "flag": worst.get(dash["id"], ""), "reason": reason.get(dash["id"], ""),
            })
        for wid in widgets:
            widget = graph.nodes[wid]
            wattrs = clean_attrs(widget["attrs"])
            hits = graph.children(wid, "filters_on_agent")
            rows.append({
                "group": group,
                "dashboard_id": attrs.get("dashboard_id"),
                "dashboard": dash["name"],
                "widget_id": wattrs.get("widget_id"),
                "widget": widget["name"],
                "widget_type": wattrs.get("widget_type") or "",
                "filters": ", ".join(wattrs.get("filter_keys") or []),
                "agents_hit": len(hits),
                "selected_agents_hit": len([h for h in hits if h in selected]),
                "flag": worst.get(wid, ""),
                "reason": reason.get(wid, ""),
            })
    return rows


def alert_rows(graph, findings):
    worst, reason = _level_info(findings)
    rows = []
    for node in sorted(graph.nodes_of("alert_rule"), key=lambda n: ((n["name"] or "").lower(), n["id"])):
        attrs = clean_attrs(node["attrs"])
        tests = graph.parents(node["id"], "alert_rule")
        agents = graph.parents(node["id"], "agent_notif_rule")
        rows.append({
            "kind": "alert rule (%s)" % (attrs.get("flavour") or ""),
            "id": attrs.get("rule_id"), "name": node["name"],
            "group": _group_names(graph, [attrs.get("aid")]),
            "detail": "severity=%s" % (attrs.get("severity") or ""),
            "tests": _names(graph, tests), "agents": _names(graph, agents),
            "flag": worst.get(node["id"], ""), "reason": reason.get(node["id"], ""),
        })
    for node in sorted(graph.nodes_of("suppression"), key=lambda n: ((n["name"] or "").lower(), n["id"])):
        attrs = clean_attrs(node["attrs"])
        rows.append({
            "kind": "suppression window",
            "id": attrs.get("suppression_id"), "name": node["name"],
            "group": _group_names(graph, [attrs.get("aid")]),
            "detail": "enabled=%s; status=%s" % (attrs.get("enabled"), attrs.get("status") or ""),
            "tests": _names(graph, graph.children(node["id"], "applies_to_test")),
            "agents": _names(graph, graph.children(node["id"], "applies_to_agent")),
            "flag": worst.get(node["id"], ""), "reason": reason.get(node["id"], ""),
        })
    return rows


def summary_rows(graph, findings, selection, meta):
    rows = [("Report", "ThousandEyes agent migration: relationship report")]
    if meta.get("generated"):
        rows.append(("Generated", meta["generated"]))
    if meta.get("input"):
        rows.append(("Input", meta["input"]))
    rows.append(("", ""))
    stats = graph.stats()
    for label, kind in (
        ("Account groups", "group"), ("Agents", "agent"), ("Clusters", "cluster"),
        ("Cluster members", "member"), ("Tests", "test"), ("Tags", "tag"),
        ("Dashboards", "dashboard"), ("Widgets", "widget"), ("Global filters", "global_filter"),
        ("Alert rules", "alert_rule"), ("Suppression windows", "suppression"),
    ):
        rows.append((label, stats.get(kind, 0)))
    rows.append(("Relationships (edges)", len(graph.edges)))
    rows.append(("", ""))
    if selection is None:
        rows.append(("Selection", "none given: inventory only, no impact flags"))
    else:
        rows.append(("Selected agents/clusters", len(selection["selected"])))
        rows.append(("Selected but offline (not migrated)", len(selection["excluded_offline"])))
        rows.append(("Unmatched identifiers", len(selection["unmatched"])))
        counts = summarize(findings or [])
        for level in LEVELS:
            rows.append(("Findings " + level, counts.get(level, 0)))
        rows.append(("Findings total", counts["total"]))
    rows.append(("", ""))
    rows.append(("OK", "no change expected"))
    rows.append(("REVIEW", "check by hand"))
    rows.append(("BREAKS", "stops working unless a fix runs"))
    return rows


def _fill(workbook, title, columns, rows, widths=None, first=False):
    sheet = workbook.active if first else workbook.create_sheet()
    sheet.title = title
    sheet.append(list(columns))
    for row in rows[:MAX_ROWS]:
        sheet.append([_cell(row.get(c)) for c in columns])
    _style(sheet, columns, widths)
    return sheet


def _style(sheet, columns, widths=None):
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="0B4A6F")
    for index, name in enumerate(columns, start=1):
        cell = sheet.cell(row=1, column=index)
        cell.font = head_font
        cell.fill = head_fill
        width = (widths or {}).get(name) or min(60, max(12, len(str(name)) + 4))
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(columns)), max(sheet.max_row, 1))


_WIDTHS = {
    "reason": 70, "fix": 50, "element_name": 40, "element_id": 28, "name": 40, "id": 28,
    "detail": 40, "agents": 40, "tests": 40, "alert_rules": 30, "tags": 30, "filters": 36,
    "dashboard": 32, "widget": 32, "groups": 30, "ips": 24, "member_ips": 24, "group_name": 28,
    "agent_name": 30,
}


def write_workbook(path, graph, findings, selection, meta):
    """Write the nine-sheet workbook to ``path`` and return ``path``.

    ``selection`` is the dict from ``resolve_selection`` (or None).
    Raises ``RuntimeError`` when openpyxl is not installed.
    """
    try:
        import openpyxl
    except ImportError:
        raise RuntimeError(MISSING_MESSAGE)

    findings = list(findings or [])
    meta = meta or {}
    selected = set(selection["selected"]) if selection else set()
    workbook = openpyxl.Workbook()

    summary = workbook.active
    summary.title = "Summary"
    summary.append(["Item", "Value"])
    for label, value in summary_rows(graph, findings, selection, meta):
        summary.append([_cell(label), _cell(value)])
    _style(summary, ["Item", "Value"], {"Item": 38, "Value": 60})

    if selection is None:
        selected_data = []
        offline_data = offline_rows(graph, all_offline_ids(graph))
    else:
        by_node = _matched_by(selection)
        selected_data = selected_rows(graph, selection)
        offline_data = offline_rows(graph, selection["excluded_offline"], by_node)
    _fill(workbook, "Selected", AGENT_COLUMNS, selected_data, _WIDTHS)
    _fill(workbook, "Offline", AGENT_COLUMNS, offline_data, _WIDTHS)
    _fill(workbook, "Shares", SHARE_COLUMNS, share_rows(graph, selected if selection else None), _WIDTHS)
    _fill(workbook, "Tree", TREE_COLUMNS, tree_rows(graph, selected if selection else None),
          dict(_WIDTHS, name=60, id=32))
    _fill(workbook, "Findings", FINDING_COLUMNS, findings, _WIDTHS)
    _fill(workbook, "Tests", TEST_COLUMNS, tests_rows(graph, findings), _WIDTHS)
    _fill(workbook, "Dashboards", DASHBOARD_COLUMNS, dashboard_rows(graph, findings, selected), _WIDTHS)
    _fill(workbook, "Alerts", ALERT_COLUMNS, alert_rows(graph, findings), _WIDTHS)
    workbook.save(path)
    return path
