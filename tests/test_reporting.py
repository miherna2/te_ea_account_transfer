import json
from pathlib import Path

from openpyxl import load_workbook

from te_agent_migrate.reporting import RunReport


def test_report_writes_redacted_json_and_required_sheets(tmp_path: Path):
    report = RunReport(tmp_path, "dry-run", {"token": "secret", "strategy": "A"})
    report.record_event({"step": "auth", "status": "ok", "account_token": "A" * 32})
    report.append_agent({"hostname": "a1"})
    report.append_decision({"prompt": "continue", "response": "yes"})
    json_path, xlsx_path = report.write()

    data = json.loads(json_path.read_text())
    assert data["metadata"]["token"] == "[REDACTED]"
    assert data["steps"][0]["account_token"] == "[REDACTED]"
    assert data["agents"][0]["hostname"] == "a1"

    wb = load_workbook(xlsx_path)
    assert wb.sheetnames == [
        "Summary",
        "Agents",
        "Tests",
        "Unsupported Tests",
        "Stale Entries",
        "Diffs",
        "Decisions",
        "Errors",
    ]
